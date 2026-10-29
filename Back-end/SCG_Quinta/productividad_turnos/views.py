from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST

from .forms import AccionForm, FiltroForm, ProgramacionForm, RegistroForm
from .models import ProgramacionTurno, RegistroProductividadTurno
from .permissions import require_permission, scoped_programs, user_label
from .selectors import metrics, programs_from_filter, registro_of, summarize
from .services import cambiar_estado, guardar_registro, programar


VIEW = "view_registroproductividadturno"


def base_context(request, **kwargs):
    return {"usuario_nombre": user_label(request.user), "inicio_url": getattr(settings, "PRODUCTIVIDAD_INICIO_URL", "/inicio/"), **kwargs}


@require_permission(VIEW)
@require_GET
def panel(request, periodo="dia"):
    form = FiltroForm(request.GET, user=request.user, periodo=periodo)
    resumen = None
    if form.is_valid():
        qs, start, end = programs_from_filter(request.user, form.cleaned_data, periodo)
        resumen = summarize(qs, start, end, periodo)
    api_query = request.GET.copy()
    api_query["periodo"] = periodo
    return render(request, "productividad_turnos/panel.html", base_context(
        request, filtro=form, resumen=resumen, periodo=periodo,
        api_url=reverse("productividad_turnos:api_resumen") + "?" + api_query.urlencode(),
        export_url=reverse("productividad_turnos:exportar") + "?" + api_query.urlencode(),
    ))


@require_permission(VIEW)
@require_GET
def api_resumen(request):
    periodo = request.GET.get("periodo", "dia")
    if periodo not in ("dia", "semana"):
        return JsonResponse({"error": "Período inválido."}, status=400)
    form = FiltroForm(request.GET, user=request.user, periodo=periodo)
    if not form.is_valid():
        return JsonResponse({"errors": form.errors.get_json_data()}, status=400)
    qs, start, end = programs_from_filter(request.user, form.cleaned_data, periodo)
    resumen = summarize(qs, start, end, periodo)
    numeric = {k: resumen[k] for k in ("total", "vencidos", "recibidos", "pendientes", "cobertura", "oportunidad", "meta_programada", "avance_plan", "actualizado")}
    numeric["html"] = render_to_string("productividad_turnos/_panel_content.html", {"resumen": resumen, "periodo": periodo}, request=request)
    return JsonResponse(numeric)


@require_permission(VIEW)
@require_GET
def historial(request):
    periodo = request.GET.get("periodo", "semana")
    if periodo not in ("dia", "semana"):
        periodo = "semana"
    form = FiltroForm(request.GET, user=request.user, periodo=periodo)
    rows = []
    if form.is_valid():
        qs, start, end = programs_from_filter(request.user, form.cleaned_data, periodo)
        from .selectors import estado_plan
        for p in qs:
            r = registro_of(p)
            rows.append({"plan": p, "registro": r, "estado": estado_plan(p), "metrics": metrics([r]) if r and r.estado != "BORRADOR" else None})
    return render(request, "productividad_turnos/historial.html", base_context(request, filtro=form, rows=rows, periodo=periodo))


@require_permission(VIEW)
@require_GET
def seleccionar(request):
    # La misma lista filtrada permite seleccionar turno sin escribir IDs ni duplicar programación.
    return historial(request)


@require_permission(VIEW)
@require_http_methods(["GET", "POST"])
def registro(request, pk):
    plan = get_object_or_404(scoped_programs(request.user), pk=pk)
    reg = registro_of(plan)
    initial = {"version": reg.version if reg else 0}
    if reg:
        initial.update({f: getattr(reg, f) for f in ("dotacion", "semielaborados", "terminadas", "horas_persona", "observaciones")})
    form = RegistroForm(request.POST if request.method == "POST" else None, initial=initial, area=plan.area)
    if request.method == "POST":
        action = request.POST.get("accion")
        if action not in ("borrador", "cerrar"):
            form.add_error(None, "Acción inválida.") if form.is_valid() else None
        elif form.is_valid():
            try:
                guardar_registro(request.user, plan.pk, form.cleaned_data, form.cleaned_data["version"], cerrar=action == "cerrar")
            except ValidationError as exc:
                form.add_error(None, " ".join(exc.messages))
            except IntegrityError:
                form.add_error(None, "Otro usuario modificó este turno. Recargue la página y revise el registro actual.")
            else:
                messages.success(request, "Turno cerrado e incorporado a los indicadores." if action == "cerrar" else "Borrador guardado. No participa de los indicadores.")
                return redirect("productividad_turnos:registro", pk=pk)
    locked = bool(reg and reg.estado != "BORRADOR")
    if locked:
        for field in form.fields.values():
            field.disabled = True
    revisions = plan.revisiones.select_related("usuario")[:30]
    return render(request, "productividad_turnos/registro.html", base_context(request, plan=plan, registro=reg, form=form, locked=locked, accion_form=AccionForm(initial={"version": reg.version if reg else 1}), revisiones=revisions))


@require_permission("verificar_turno")
@require_POST
def verificar(request, pk):
    return _action(request, pk, "verificar")


@require_permission("reabrir_turno")
@require_POST
def reabrir(request, pk):
    return _action(request, pk, "reabrir")


def _action(request, pk, action):
    get_object_or_404(scoped_programs(request.user), pk=pk)
    form = AccionForm(request.POST)
    if not form.is_valid():
        messages.error(request, "Solicitud inválida. Recargue la página y complete el motivo cuando corresponda.")
    else:
        try:
            cambiar_estado(request.user, pk, action, form.cleaned_data["version"], form.cleaned_data["motivo"])
        except ValidationError as exc:
            messages.error(request, " ".join(exc.messages))
        except RegistroProductividadTurno.DoesNotExist:
            messages.error(request, "El turno no tiene registro.")
        else:
            messages.success(request, "Cierre verificado." if action == "verificar" else "Turno reabierto. Queda fuera de los indicadores hasta cerrarlo nuevamente.")
    return redirect("productividad_turnos:registro", pk=pk)


@require_permission("programar_turno")
@require_http_methods(["GET", "POST"])
def programacion(request, pk=None):
    plan = get_object_or_404(scoped_programs(request.user), pk=pk) if pk else None
    initial = {"fecha_operativa": timezone.localdate(), "version": 0, "programado": True}
    if plan:
        initial.update({f: getattr(plan, f) for f in ("fecha_operativa", "planta", "area", "turno", "meta_pt", "programado", "extraordinario", "motivo", "version")})
    form = ProgramacionForm(request.POST if request.method == "POST" else None, user=request.user, initial=initial, editing=bool(plan))
    if request.method == "POST" and form.is_valid():
        try:
            plan = programar(request.user, form.cleaned_data, pk, form.cleaned_data["version"])
        except ValidationError as exc:
            form.add_error(None, " ".join(exc.messages))
        except IntegrityError:
            form.add_error(None, "Ya existe una programación con esa fecha, planta, área y turno. Use Editar desde el historial.")
        else:
            messages.success(request, "Programación guardada.")
            return redirect("productividad_turnos:registro", pk=plan.pk)
    return render(request, "productividad_turnos/programacion.html", base_context(request, form=form, plan=plan))


@require_permission(VIEW)
@require_GET
def exportar(request):
    periodo = request.GET.get("periodo", "dia")
    if periodo not in ("dia", "semana"):
        return HttpResponse("Período inválido", status=400)
    form = FiltroForm(request.GET, user=request.user, periodo=periodo)
    if not form.is_valid():
        return HttpResponse("Filtros inválidos", status=400)
    plans, start, end = programs_from_filter(request.user, form.cleaned_data, periodo)
    # openpyxl ya forma parte de requirements.txt. La importación local mantiene
    # independiente la carga de las demás páginas de la aplicación.
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .selectors import estado_plan

    workbook = Workbook()
    workbook.properties.creator = "SCG_Quinta"
    sheet = workbook.active
    sheet.title = "Turnos"
    sheet.sheet_view.showGridLines = False
    headers = ["Fecha operativa", "Planta", "Área", "Turno", "Estado", "Dotación", "Semielaborados", "Terminadas", "Meta programada", "Meta congelada cierre", "Horas-persona", "Cumplimiento", "Unidades/persona-turno", "Unidades/HH", "Responsable cierre", "Fecha cierre", "Observaciones"]
    metadata = [
        "Quinta · Productividad por turnos",
        "Panel {0}: {1:%d/%m/%Y} al {2:%d/%m/%Y}".format("semanal" if periodo == "semana" else "diario", start, end),
        "Fechas de cierre en zona horaria: {0}".format(timezone.get_default_timezone()),
        "Incluye los estados de los turnos filtrados. Los indicadores solo corresponden a cierres de turnos programados. Los semielaborados se registran por separado.",
    ]
    for index, text in enumerate(metadata, 1):
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=len(headers))
        cell = sheet.cell(index, 1, text)
        cell.font = Font(name="Calibri", size=16 if index == 1 else 11, bold=index == 1, color="FFFFFF" if index == 1 else "173C35")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if index == 1:
            cell.fill = PatternFill("solid", fgColor="173C35")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[4].height = 32
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(5, column, label)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="206B54")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[5].height = 34

    for row_index, p in enumerate(plans, 6):
        r = registro_of(p)
        official = r and r.estado != "BORRADOR" and p.programado
        m = metrics([r]) if official else {}
        closed_at = timezone.localtime(r.cerrado, timezone.get_default_timezone()).replace(tzinfo=None) if official else None
        compliance = m.get("cumplimiento")
        vals = [p.fecha_operativa, p.get_planta_display(), p.area.nombre, p.turno, estado_plan(p), r.dotacion if r else None, r.semielaborados if r else None, r.terminadas if r else None, p.meta_pt, r.meta_pt_cierre if official else None, r.horas_persona if r else None, compliance / Decimal(100) if compliance is not None else None, m.get("up"), m.get("uph"), user_label(r.cerrado_por) if official and r.cerrado_por else None, closed_at, r.observaciones if r else None]
        for column, value in enumerate(vals, 1):
            cell = sheet.cell(row_index, column)
            if isinstance(value, str):
                # Guardar siempre texto literal, incluso si comienza con '='.
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                cell.data_type = "s"
            else:
                cell.value = value
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=column == 17)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5F2")
            if column == 1:
                cell.number_format = "dd/mm/yyyy"
            elif column == 16:
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            elif 6 <= column <= 10:
                cell.number_format = "#,##0"
            elif column == 12:
                cell.number_format = "0.00%"
            elif column in (11, 13, 14):
                cell.number_format = "#,##0.00"

    widths = [18, 23, 22, 10, 20, 13, 18, 16, 19, 21, 18, 18, 24, 18, 30, 24, 60]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = "A5:Q{0}".format(max(5, sheet.max_row))
    sheet.print_title_rows = "1:5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="productividad_{start}_{end}.xlsx"'
    return response