from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import patch
try:
    from zoneinfo import ZoneInfo
except ImportError:
    from backports.zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError, transaction
from django.test import Client, TestCase
from django.urls import reverse

from productividad_turnos.models import AreaProductividad, AsignacionProductividad, ProgramacionTurno, RegistroProductividadTurno, RevisionProductividad
from productividad_turnos.selectors import metrics, programs_from_filter, scoped_programs, summarize
from productividad_turnos.services import cambiar_estado, guardar_registro, horario, programar
from productividad_turnos.forms import FiltroForm, RegistroForm


class ProductividadTests(TestCase):
    def setUp(self):
        self.now = datetime(2026, 10, 15, 8, tzinfo=ZoneInfo("America/Santiago"))
        self.clock = patch("django.utils.timezone.now", return_value=self.now)
        self.clock.start()
        self.addCleanup(self.clock.stop)
        self.area, _ = AreaProductividad.objects.get_or_create(codigo="TORTAS", defaults={"nombre": "Tortas"})
        self.user = self.make_user("supervisor")
        self.user.user_permissions.add(*Permission.objects.filter(content_type__app_label="productividad_turnos"))
        AsignacionProductividad.objects.create(usuario=self.user, planta="ENEA", area=self.area)

    def make_user(self, name):
        User = get_user_model()
        if User.USERNAME_FIELD == "nombre_completo":
            user = User.objects.create_user(nombre_completo=name, perfil_usuario="Supervisor", rut=f"{get_user_model().objects.count()+1}-9", password="test-password")
        else:
            user = User.objects.create_user(username=name, password="test-password")
        user.is_active = True
        user.save()
        return user

    def plan_data(self, fecha=date(2026, 10, 13), turno="A", meta=3800, planta="ENEA"):
        return {"fecha_operativa": fecha, "planta": planta, "area": self.area, "turno": turno, "meta_pt": meta, "programado": True, "extraordinario": False, "motivo": "Prueba"}

    def plan(self, **kwargs):
        return programar(self.user, self.plan_data(**kwargs))

    def values(self, pt=2975, dot=53, semi=3464, hh=None):
        return {"dotacion": dot, "semielaborados": semi, "terminadas": pt, "horas_persona": hh, "observaciones": "Prueba", "confirmar_ceros": True}

    def close(self, plan, **kwargs):
        return guardar_registro(self.user, plan.pk, self.values(**kwargs), 0, cerrar=True)

    def query(self, **overrides):
        data = {"fecha": "2026-10-13", "anio": "2026", "semana": "42", "planta": "ENEA", "area": str(self.area.pk), "turnos": ["A", "B", "C"]}
        data.update(overrides)
        return data

    def excel_sheet(self, response):
        from openpyxl import load_workbook
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))["Turnos"]

    def example(self):
        for fecha, rows in [(date(2026, 10, 13), [("A", 53, 3464, 2975, 3800), ("B", 82, 4652, 5599, 5700), ("C", 65, 3291, 4492, 5500)]), (date(2026, 10, 14), [("A", 50, 3488, 2446, 3800), ("B", 82, 4300, 4813, 5700), ("C", 64, 3767, 4609, 5500)])]:
            for t, dot, semi, pt, meta in rows:
                self.close(self.plan(fecha=fecha, turno=t, meta=meta), pt=pt, dot=dot, semi=semi)

    def test_real_excel_totals_and_weighted_productivity(self):
        self.example()
        rows = RegistroProductividadTurno.objects.all()
        tuesday = metrics(rows.filter(programacion__fecha_operativa=date(2026, 10, 13)))
        self.assertEqual(tuesday["terminadas"], 13066)
        self.assertEqual(tuesday["up"], Decimal("65.33"))
        self.assertEqual(metrics(rows.filter(programacion__fecha_operativa=date(2026, 10, 14)))["semielaborados"], 11555)
        total = metrics(rows)
        self.assertEqual(total["terminadas"], 24934)
        self.assertEqual(total["personas_turno"], 396)
        self.assertEqual(round(total["cumplimiento"], 2), Decimal("83.11"))
        self.assertEqual(round(total["up"], 2), Decimal("62.96"))

    def test_unequal_targets_are_weighted(self):
        rows = [SimpleNamespace(terminadas=50, meta_pt_cierre=100, semielaborados=0, dotacion=1, horas_persona=None), SimpleNamespace(terminadas=900, meta_pt_cierre=1000, semielaborados=0, dotacion=1, horas_persona=None)]
        self.assertEqual(round(metrics(rows)["cumplimiento"], 2), Decimal("86.36"))

    def test_empty_missing_zero_and_hours(self):
        self.assertIsNone(metrics([])["cumplimiento"])
        r = SimpleNamespace(terminadas=0, meta_pt_cierre=100, semielaborados=0, dotacion=2, horas_persona=None)
        self.assertEqual(metrics([r])["cumplimiento"], 0)
        r.terminadas, r.meta_pt_cierre = 10, 0
        self.assertIsNone(metrics([r])["cumplimiento"])
        self.assertTrue(metrics([r])["metas_incompletas"])
        r.dotacion = None
        self.assertIsNone(metrics([r])["up"])
        self.assertIsNone(metrics([r])["uph"])

    def test_above_one_hundred_not_capped(self):
        r = self.close(self.plan(meta=100), pt=250)
        self.assertEqual(metrics([r])["cumplimiento"], 250)

    def test_night_shift_and_iso_year(self):
        start, end = horario(date(2026, 10, 26), "A")
        self.assertEqual(start.date(), date(2026, 10, 25))
        self.assertEqual(end.date(), date(2026, 10, 26))
        self.assertEqual(date(2027, 1, 1).isocalendar()[0], 2026)
        self.assertTrue(FiltroForm(self.query(anio="2026", semana="53"), user=self.user, periodo="semana").is_valid())
        self.assertFalse(FiltroForm(self.query(anio="2025", semana="53"), user=self.user, periodo="semana").is_valid())

    def test_schedule_draft_close_verify_reopen_audit(self):
        p = self.plan()
        draft = guardar_registro(self.user, p.pk, self.values(), 0)
        self.assertEqual(draft.estado, "BORRADOR")
        closed = guardar_registro(self.user, p.pk, self.values(), draft.version, cerrar=True)
        self.assertEqual(closed.meta_pt_cierre, 3800)
        checked = cambiar_estado(self.user, p.pk, "verificar", closed.version)
        self.assertEqual(checked.estado, "VERIFICADO")
        with self.assertRaises(ValidationError):
            cambiar_estado(self.user, p.pk, "reabrir", checked.version)
        opened = cambiar_estado(self.user, p.pk, "reabrir", checked.version, "Corrección dotación")
        self.assertEqual(opened.estado, "BORRADOR")
        self.assertIsNone(opened.cerrado)
        self.assertTrue(RevisionProductividad.objects.filter(programacion=p, accion="REABRIR", motivo="Corrección dotación").exists())

    def test_stale_version_does_not_overwrite(self):
        p = self.plan()
        guardar_registro(self.user, p.pk, self.values(), 0)
        with self.assertRaises(ValidationError):
            guardar_registro(self.user, p.pk, self.values(pt=999), 0)
        self.assertEqual(RegistroProductividadTurno.objects.get(programacion=p).terminadas, 2975)

    def test_closed_record_is_immutable(self):
        p = self.plan()
        r = self.close(p)
        with self.assertRaises(ValidationError):
            guardar_registro(self.user, p.pk, self.values(), r.version)

    def test_duplicate_program_and_record_protected_by_database(self):
        p = self.plan()
        with self.assertRaises(ValidationError):
            self.plan()
        guardar_registro(self.user, p.pk, self.values(), 0)
        with self.assertRaises(IntegrityError), transaction.atomic():
            RegistroProductividadTurno.objects.create(programacion=p, creado_por=self.user)

    def test_future_turn_cannot_close_but_can_save_draft(self):
        p = self.plan(fecha=date(2026, 10, 15), turno="B")
        with self.assertRaises(ValidationError):
            self.close(p)
        r = guardar_registro(self.user, p.pk, self.values(), 0)
        self.assertEqual(r.estado, "BORRADOR")

    def test_close_requires_people_and_confirmed_zero(self):
        p = self.plan()
        with self.assertRaises(ValidationError):
            self.close(p, dot=None)
        data = self.values(pt=0)
        data["confirmar_ceros"] = False
        with self.assertRaises(ValidationError):
            guardar_registro(self.user, p.pk, data, 0, cerrar=True)
        data["confirmar_ceros"], data["observaciones"] = True, ""
        with self.assertRaises(ValidationError):
            guardar_registro(self.user, p.pk, data, 0, cerrar=True)

    def test_optional_hours_not_estimated(self):
        a = self.close(self.plan(), hh=Decimal("100"))
        b = self.close(self.plan(turno="B"), hh=None)
        self.assertEqual(metrics([a])["uph"], Decimal("29.75"))
        self.assertIsNone(metrics([a, b])["uph"])
        f = RegistroForm({"version": "0", "horas_persona": "0.01"}, area=self.area)
        self.assertTrue(f.is_valid(), f.errors)

    def test_frozen_target_and_reprogramming(self):
        p = self.plan()
        r = self.close(p)
        data = self.plan_data(meta=4000)
        with self.assertRaises(ValidationError):
            programar(self.user, data, p.pk, p.version)
        opened = cambiar_estado(self.user, p.pk, "reabrir", r.version, "Cambio autorizado")
        updated = programar(self.user, data, p.pk, p.version)
        closed = guardar_registro(self.user, p.pk, self.values(), opened.version, cerrar=True)
        self.assertEqual(closed.meta_pt_cierre, 4000)
        self.assertEqual(updated.version, 2)
        self.assertEqual(RevisionProductividad.objects.filter(accion="CIERRE").first().antes.get("meta_pt_cierre"), 3800)

    def test_zero_target_requires_reason(self):
        d = self.plan_data(meta=0)
        d["motivo"] = ""
        with self.assertRaises(ValidationError):
            programar(self.user, d)

    def test_unprogrammed_or_inactive_cannot_receive_record(self):
        d = self.plan_data()
        d["programado"] = False
        p = programar(self.user, d)
        with self.assertRaises(ValidationError):
            self.close(p)
        self.area.activo = False
        self.area.save()
        self.assertTrue(FiltroForm(self.query(), user=self.user).is_valid())

    def test_coverage_excludes_future_and_unprogrammed(self):
        self.close(self.plan())
        self.plan(turno="B")
        self.plan(fecha=date(2026, 10, 15), turno="B")
        d = self.plan_data(turno="C")
        d["programado"] = False
        programar(self.user, d)
        report = summarize(scoped_programs(self.user), date(2026, 10, 12), date(2026, 10, 18), "semana", now=self.now)
        self.assertEqual(report["vencidos"], 2)
        self.assertEqual(report["pendientes"], 1)
        self.assertEqual(report["cobertura"], 50)
        self.assertEqual(report["total"]["cierres"], 1)

    def test_scope_service_denies_other_plant(self):
        with self.assertRaises(PermissionDenied):
            self.plan(planta="PUERTO_VESPUCIO")

    def test_panel_api_export_and_templates(self):
        self.example()
        self.client.force_login(self.user)
        for name in ("diario", "semanal", "historial", "programacion"):
            response = self.client.get(reverse(f"productividad_turnos:{name}"), self.query())
            self.assertEqual(response.status_code, 200, response.content[:300])
        p = ProgramacionTurno.objects.first()
        self.assertEqual(self.client.get(reverse("productividad_turnos:registro", args=[p.pk])).status_code, 200)
        api = self.client.get(reverse("productividad_turnos:api_resumen"), self.query(periodo="semana"))
        self.assertEqual(api.status_code, 200)
        self.assertEqual(api.json()["total"]["terminadas"], 24934)
        self.assertIn("13.066", api.json()["html"])
        export = self.client.get(reverse("productividad_turnos:exportar"), self.query(periodo="semana"))
        self.assertEqual(export.status_code, 200)
        self.assertEqual(export["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertTrue(export["Content-Disposition"].endswith('.xlsx"'))
        self.assertEqual(len(list(self.excel_sheet(export).iter_rows(min_row=6))), 6)

    def test_invalid_filter_or_empty_turns_cannot_broaden_scope(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("productividad_turnos:api_resumen"), self.query(area="99999"))
        self.assertEqual(response.status_code, 400)
        response = self.client.get(reverse("productividad_turnos:api_resumen"), self.query(turnos=[], filtrar="1"))
        self.assertEqual(response.status_code, 400)

    def test_http_permissions_and_read_only_actions(self):
        p = self.plan()
        anonymous = Client().get(reverse("productividad_turnos:diario"))
        self.assertEqual(anonymous.status_code, 302)
        self.user.user_permissions.clear()
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("productividad_turnos:diario")).status_code, 403)

    def test_scope_view_and_export_do_not_leak(self):
        p = self.plan()
        ProgramacionTurno.objects.filter(pk=p.pk).update(planta="PUERTO_VESPUCIO")
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("productividad_turnos:registro", args=[p.pk])).status_code, 404)
        response = self.client.get(reverse("productividad_turnos:api_resumen"), self.query(planta=""))
        self.assertEqual(response.json()["total"]["cierres"], 0)
        response = self.client.get(reverse("productividad_turnos:exportar"), self.query(planta=""))
        rows = list(self.excel_sheet(response).iter_rows(min_row=6, values_only=True))
        self.assertEqual(rows, [])

    def test_csrf_enforced_for_writes(self):
        p = self.plan()
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(reverse("productividad_turnos:registro", args=[p.pk]), {"accion": "cerrar"}).status_code, 403)

    def test_edit_via_http_and_stale_form_message(self):
        p = self.plan()
        self.client.force_login(self.user)
        url = reverse("productividad_turnos:registro", args=[p.pk])
        payload = {"version": 0, "dotacion": 53, "semielaborados": 3464, "terminadas": 2975, "accion": "cerrar"}
        self.assertEqual(self.client.post(url, payload).status_code, 302)
        response = self.client.post(url, payload)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "El registro cambió")
        self.assertEqual(RegistroProductividadTurno.objects.count(), 1)

    def test_history_filter_preserves_period(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("productividad_turnos:historial"), self.query(periodo="dia"))
        self.assertContains(response, 'name="periodo" value="dia"')

    def test_command_bootstrap_is_idempotent(self):
        call_command("configurar_productividad", stdout=StringIO())
        call_command("configurar_productividad", usuario_id=self.user.pk, grupo="Supervisor", planta="ENEA", stdout=StringIO())
        self.assertEqual(AsignacionProductividad.objects.filter(usuario=self.user).count(), 1)
        self.assertTrue(self.user.groups.filter(name="Productividad - Supervisor").exists())

    def test_week_generation_rolls_back_duplicates(self):
        self.plan(fecha=date(2026, 10, 13))
        before = ProgramacionTurno.objects.count()
        with self.assertRaises(CommandError):
            call_command("programar_semana_productividad", usuario_id=self.user.pk, planta="ENEA", anio=2026, semana=42, dias=[1, 2], turnos=["A"], meta_a=3800, stdout=StringIO())
        self.assertEqual(ProgramacionTurno.objects.count(), before)

    def test_excel_formula_and_template_escaping(self):
        p = self.plan()
        self.close(p)
        RegistroProductividadTurno.objects.filter(programacion=p).update(observaciones="=SUM(1+1)")
        self.client.force_login(self.user)
        response = self.client.get(reverse("productividad_turnos:exportar"), self.query())
        cell = self.excel_sheet(response)["Q6"]
        self.assertEqual(cell.value, "=SUM(1+1)")
        self.assertEqual(cell.data_type, "s")
        RevisionProductividad.objects.create(programacion=p, usuario=self.user, accion="TEST", motivo="<script>alert(1)</script>")
        response = self.client.get(reverse("productividad_turnos:registro", args=[p.pk]))
        self.assertContains(response, "&lt;script&gt;")

    def test_excel_numeric_dates_semifinished_and_percent(self):
        p = self.plan()
        self.close(p, hh=Decimal("424.50"))
        self.client.force_login(self.user)
        response = self.client.get(reverse("productividad_turnos:exportar"), self.query())
        sheet = self.excel_sheet(response)
        self.assertEqual(sheet["A6"].value, datetime(2026, 10, 13))
        self.assertEqual(sheet["F6"].value, 53)
        self.assertEqual(sheet["G6"].value, 3464)
        self.assertEqual(sheet["H6"].value, 2975)
        self.assertEqual(sheet["I6"].value, 3800)
        self.assertEqual(sheet["J6"].value, 3800)
        self.assertEqual(sheet["K6"].value, 424.5)
        self.assertAlmostEqual(sheet["L6"].value, 2975 / 3800)
        self.assertEqual(sheet["L6"].number_format, "0.00%")
        self.assertAlmostEqual(sheet["M6"].value, 2975 / 53)
        self.assertAlmostEqual(sheet["N6"].value, 2975 / 424.5)
        self.assertEqual(sheet["P6"].value, datetime(2026, 10, 15, 8))
        self.assertIsNone(sheet["P6"].value.tzinfo)
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertEqual(sheet.auto_filter.ref, "A5:Q6")
        self.assertContains(self.client.get(reverse("productividad_turnos:diario"), self.query()), "Descargar Excel")

    def test_excel_filters_week_day_and_shift(self):
        self.example()
        self.close(self.plan(fecha=date(2026, 10, 5)), pt=999)
        self.client.force_login(self.user)
        url = reverse("productividad_turnos:exportar")
        day = self.excel_sheet(self.client.get(url, self.query(turnos=["B"])))
        rows = list(day.iter_rows(min_row=6, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][3], "B")
        self.assertEqual(rows[0][7], 5599)
        week = self.excel_sheet(self.client.get(url, self.query(periodo="semana", turnos=["A"])))
        self.assertEqual([r[7] for r in week.iter_rows(min_row=6, values_only=True)], [2975, 2446])

    def test_excel_drafts_and_missing_values_do_not_become_official_zeros(self):
        self.plan(turno="A")
        draft = self.plan(turno="B")
        guardar_registro(self.user, draft.pk, self.values(pt=22, semi=12), 0)
        self.client.force_login(self.user)
        sheet = self.excel_sheet(self.client.get(reverse("productividad_turnos:exportar"), self.query()))
        self.assertEqual(sheet["E6"].value, "Pendiente")
        self.assertIsNone(sheet["H6"].value)
        self.assertEqual(sheet["E7"].value, "Borrador")
        self.assertEqual(sheet["G7"].value, 12)
        self.assertEqual(sheet["H7"].value, 22)
        for row in (6, 7):
            for column in (10, 12, 13, 14, 16):
                self.assertIsNone(sheet.cell(row, column).value)

    def test_excel_zero_above_one_hundred_and_missing_meta(self):
        self.close(self.plan(turno="A", meta=100), pt=250)
        self.close(self.plan(turno="B", meta=100), pt=0, semi=0)
        self.close(self.plan(turno="C", meta=0), pt=50)
        self.client.force_login(self.user)
        sheet = self.excel_sheet(self.client.get(reverse("productividad_turnos:exportar"), self.query()))
        self.assertEqual(sheet["L6"].value, 2.5)
        self.assertEqual(sheet["L7"].value, 0)
        self.assertEqual(sheet["G7"].value, 0)
        self.assertIsNone(sheet["L8"].value)

    def test_excel_empty_invalid_and_access_controls(self):
        url = reverse("productividad_turnos:exportar")
        self.assertEqual(self.client.get(url, self.query()).status_code, 302)
        self.client.force_login(self.user)
        empty = self.client.get(url, self.query())
        self.assertEqual(self.excel_sheet(empty).max_row, 5)
        self.assertIn("productividad_2026-10-13_2026-10-13.xlsx", empty["Content-Disposition"])
        self.assertEqual(ProgramacionTurno.objects.count(), 0)
        for query in (self.query(periodo="otro"), self.query(area="99999"), self.query(turnos=[], filtrar="1")):
            self.assertEqual(self.client.get(url, query).status_code, 400)
        self.assertEqual(self.client.post(url, self.query()).status_code, 405)
        self.user.user_permissions.clear()
        self.assertEqual(self.client.get(url, self.query()).status_code, 403)

    def test_excel_free_text_is_literal_and_control_characters_are_cleaned(self):
        p = self.plan()
        self.close(p)
        self.area.nombre = "=1+1"
        self.area.save()
        self.client.force_login(self.user)
        for text in ("=HYPERLINK(\"https://example.com\")", "+SUM(1,2)", "-1+2", "@SUM(1,2)", "Texto\x01 con control", "Línea 1\nLínea 2"):
            RegistroProductividadTurno.objects.filter(programacion=p).update(observaciones=text)
            sheet = self.excel_sheet(self.client.get(reverse("productividad_turnos:exportar"), self.query()))
            self.assertEqual(sheet["C6"].value, "=1+1")
            self.assertEqual(sheet["C6"].data_type, "s")
            self.assertEqual(sheet["Q6"].value, text.replace("\x01", ""))
            self.assertEqual(sheet["Q6"].data_type, "s")

    def test_non_applicable_semifinished_is_not_zero(self):
        self.area.mide_semielaborado = False
        self.area.save()
        p = self.plan()
        r = self.close(p, semi=None)
        self.assertIsNone(r.semielaborados)
        self.assertIsNone(metrics([r])["semielaborados"])

    def test_database_rejects_negative_hours(self):
        r = guardar_registro(self.user, self.plan().pk, self.values(), 0)
        with self.assertRaises(IntegrityError), transaction.atomic():
            RegistroProductividadTurno.objects.filter(pk=r.pk).update(horas_persona=Decimal("-1"))

    def test_history_with_pending_draft_and_non_programmed(self):
        self.plan()
        p = self.plan(turno="B")
        guardar_registro(self.user, p.pk, {"observaciones": "Pendiente"}, 0)
        d = self.plan_data(turno="C")
        d["programado"] = False
        programar(self.user, d)
        self.client.force_login(self.user)
        response = self.client.get(reverse("productividad_turnos:historial"), self.query())
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Pendiente")
        self.assertContains(response, "Borrador")
        self.assertContains(response, "No programado")

    def test_scope_area_pair_not_cross_product(self):
        second = AreaProductividad.objects.create(codigo="OTRA", nombre="Otra área")
        AsignacionProductividad.objects.create(usuario=self.user, planta="PUERTO_VESPUCIO", area=second)
        with self.assertRaises(PermissionDenied):
            self.plan(planta="PUERTO_VESPUCIO")

    def test_verification_and_reopening_http(self):
        p = self.plan()
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("productividad_turnos:verificar", args=[p.pk])).status_code, 405)
        self.assertEqual(self.client.post(reverse("productividad_turnos:verificar", args=[p.pk]), {"version": 1}).status_code, 302)
        r = self.close(p)
        self.client.post(reverse("productividad_turnos:verificar", args=[p.pk]), {"version": r.version})
        r.refresh_from_db()
        self.assertEqual(r.estado, "VERIFICADO")
        self.client.post(reverse("productividad_turnos:reabrir", args=[p.pk]), {"version": r.version, "motivo": "Corrección"})
        r.refresh_from_db()
        self.assertEqual(r.estado, "BORRADOR")

    def test_other_author_draft_requires_reviewer(self):
        p = self.plan()
        r = guardar_registro(self.user, p.pk, self.values(), 0)
        other = self.make_user("otro supervisor")
        other.user_permissions.add(*Permission.objects.filter(content_type__app_label="productividad_turnos", codename__in=["view_registroproductividadturno", "registrar_turno", "cerrar_turno"]))
        AsignacionProductividad.objects.create(usuario=other, planta="ENEA", area=self.area)
        with self.assertRaises(PermissionDenied):
            guardar_registro(other, p.pk, self.values(), r.version)

    def test_http_duplicate_program_is_handled(self):
        self.plan()
        self.client.force_login(self.user)
        payload = {"fecha_operativa": "2026-10-13", "planta": "ENEA", "area": self.area.pk, "turno": "A", "meta_pt": "3800", "programado": "on", "version": "0"}
        response = self.client.post(reverse("productividad_turnos:programacion"), payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ProgramacionTurno.objects.count(), 1)

    def test_chilean_number_rounding(self):
        from productividad_turnos.templatetags.productividad import numero
        self.assertEqual(numero(Decimal("19352.969"), 2), "19.352,97")
        self.assertEqual(numero(None), "—")